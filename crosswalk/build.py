"""Build the dashboard bundle and export files from the YAML source data."""
from __future__ import annotations

import csv
import json
from datetime import date
from pathlib import Path

from . import __version__
from .analysis import coverage, risk_heatmap
from .model import ROOT, Crosswalk

TEMPLATE = ROOT / "dashboard" / "template.html"
DOCS_DIR = ROOT / "docs"
EXPORT_DIR = ROOT / "exports"
PLACEHOLDER = "/*__CROSSWALK_DATA__*/null"


def bundle(cw: Crosswalk) -> dict:
    """Everything the dashboard needs, as plain JSON."""
    return {
        "meta": {
            "dataset_version": cw.version,
            "tool_version": __version__,
            "built": date.today().isoformat(),
        },
        "framework_order": cw.framework_ids,
        "frameworks": {
            fw_id: {k: fw[k] for k in ("id", "name", "short", "kind", "version", "source_url", "licence_note", "items") if k in fw}
            for fw_id, fw in cw.frameworks.items()
        },
        "mitigation_categories": list(cw.mitigation_categories.values()),
        "sectors": cw.sectors,
        "controls": cw.controls,
        "coverage": coverage(cw),
        "heatmap": risk_heatmap(cw),
    }


def build_dashboard(cw: Crosswalk, out_dir: Path = DOCS_DIR) -> Path:
    """Write docs/data.json and a self-contained docs/index.html."""
    out_dir.mkdir(parents=True, exist_ok=True)
    data = bundle(cw)
    (out_dir / "data.json").write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")

    template = TEMPLATE.read_text(encoding="utf-8")
    if PLACEHOLDER not in template:
        raise RuntimeError(f"Placeholder {PLACEHOLDER!r} not found in {TEMPLATE}")
    # Escape '</' so data can never close the <script> tag early.
    inline = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    html = template.replace(PLACEHOLDER, inline)
    out = out_dir / "index.html"
    out.write_text(html, encoding="utf-8")
    (out_dir / ".nojekyll").write_text("", encoding="utf-8")
    return out


def _long_rows(cw: Crosswalk):
    for c in cw.controls:
        for fw_id in cw.framework_ids:
            for iid in c["maps"].get(fw_id, []):
                item = cw.item(fw_id, iid)
                yield {
                    "control_id": c["id"],
                    "control_title": c["title"],
                    "mitigation_category": cw.mitigation_categories[c["category"]]["title"],
                    "framework": cw.frameworks[fw_id]["short"],
                    "item_id": iid,
                    "item_title": item["title"],
                    **{f"priority_{s}": c["sectors"].get(s, "") for s in cw.sector_ids},
                }


def export_csv(cw: Crosswalk, out_dir: Path = EXPORT_DIR) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    rows = list(_long_rows(cw))
    path = out_dir / "crosswalk_long.csv"
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    return path


def export_xlsx(cw: Crosswalk, out_dir: Path = EXPORT_DIR) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    out_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    wrap = Alignment(wrap_text=True, vertical="top")

    def write_sheet(ws, headers, rows, widths):
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill, cell.alignment = header_font, header_fill, wrap
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions

    # 1. Matrix: one row per control, one column per framework
    ws = wb.active
    ws.title = "Crosswalk matrix"
    fw_ids = cw.framework_ids
    headers = ["Control", "Title", "Mitigation category"] + [cw.frameworks[f]["short"] for f in fw_ids] + \
              [s["name"] for s in cw.sectors]
    rows = [
        [c["id"], c["title"], cw.mitigation_categories[c["category"]]["title"]]
        + [", ".join(c["maps"].get(f, [])) for f in fw_ids]
        + [c["sectors"].get(s, "") for s in cw.sector_ids]
        for c in cw.controls
    ]
    write_sheet(ws, headers, rows, [11, 38, 26] + [22] * len(fw_ids) + [14] * len(cw.sector_ids))

    # 2. Control detail
    ws = wb.create_sheet("Control detail")
    rows = [
        [c["id"], c["title"], c["objective"], "\n".join(c.get("guidance", [])),
         "\n".join(c.get("evidence", [])),
         "\n".join(f"{k}: {v}" for k, v in c.get("sector_notes", {}).items())]
        for c in cw.controls
    ]
    write_sheet(ws, ["Control", "Title", "Objective", "Implementation guidance", "Evidence", "Sector notes"],
                rows, [11, 32, 50, 60, 40, 45])

    # 3. Long mapping
    ws = wb.create_sheet("Long mapping")
    long_rows = list(_long_rows(cw))
    write_sheet(ws, list(long_rows[0].keys()), [list(r.values()) for r in long_rows],
                [11, 36, 26, 14, 12, 50] + [14] * len(cw.sector_ids))

    # 4. Coverage summary
    ws = wb.create_sheet("Coverage")
    cov = coverage(cw)
    rows = [[cw.frameworks[f]["short"], d["total"], d["covered"], d["percent"], ", ".join(d["gaps"])]
            for f, d in cov.items()]
    write_sheet(ws, ["Framework", "Items", "Covered", "Coverage %", "Uncovered items"], rows, [18, 8, 10, 12, 60])

    # 5. Sources
    ws = wb.create_sheet("Sources")
    rows = [[fw["name"], fw.get("version", ""), fw["source_url"], fw.get("licence_note", "").strip()]
            for fw in (cw.frameworks[f] for f in fw_ids)]
    write_sheet(ws, ["Framework", "Version", "Source", "Notes"], rows, [45, 30, 45, 80])

    path = out_dir / "ai_control_crosswalk.xlsx"
    wb.save(path)
    return path


def export_markdown(cw: Crosswalk, out_dir: Path = EXPORT_DIR) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    fw_ids = cw.framework_ids
    lines = ["# AI Control Crosswalk", "",
             f"Dataset version {cw.version}. Generated by `python -m crosswalk build` - do not edit by hand.", "",
             "| Control | " + " | ".join(cw.frameworks[f]["short"] for f in fw_ids) + " |",
             "|---|" + "---|" * len(fw_ids)]
    for c in cw.controls:
        cells = [", ".join(c["maps"].get(f, [])) or "-" for f in fw_ids]
        lines.append(f"| **{c['id']}** {c['title']} | " + " | ".join(cells) + " |")
    path = out_dir / "CROSSWALK.md"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def build_all(cw: Crosswalk) -> list[Path]:
    return [build_dashboard(cw), export_csv(cw), export_xlsx(cw), export_markdown(cw)]
